"""Create an evidence-aware review workbook from the existing YouTube library."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent
OUTPUT_ROOT = ROOT.parents[1] / "TOOL Video Flow OUTPUT" / "Music Library"
OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
SOURCE = ROOT / "YouTube_Library_600_Tracks.xlsx"
AUDIT = OUTPUT_ROOT / "library_audit.json"
if not AUDIT.exists():
    AUDIT = ROOT / "library_audit.json"
DEST = OUTPUT_ROOT / "YouTube_Library_Review_3_Channels.xlsx"

POOLS = {
    "01_CHRISTMAS": ("Chrismas", "Kênh 1 – HEARTHLUNE", "Jazz Noel/piano/acoustic; nghe xác nhận nhịp và bản instrumental."),
    "02_WINTER_COZY": ("Winter_Cozy", "Kênh 1 – HEARTHLUNE; Kênh 2 – SYLVANREST", "Không khí mùa đông/ấm áp; cần nghe kiểm tra tempo và chất nhạc."),
    "03_CAFE_JAZZ": ("Cafe_Jazz", "Kênh 1 – HEARTHLUNE", "Cafe jazz phù hợp định vị; loại các bản swing/nhịp mạnh sau khi nghe."),
    "04_STUDY_READING": ("Study_Reading", "Kênh 2 – SYLVANREST; Kênh 3 – ASTRALOOM", "Chỉ giữ bản nền không lời, ít tiết tấu; ưu tiên kiểm tra giai điệu nổi bật."),
    "05_DEEP_RELAX": ("Deep_Relax", "Kênh 2 – SYLVANREST; Kênh 3 – ASTRALOOM", "Ứng viên ambient/nhẹ; Kênh 3 chỉ dùng nếu không có nhịp/giai điệu rõ."),
    "06_ROMANTIC_DINNER": ("Romantic_Dinner", "Kênh 1 – HEARTHLUNE (chọn lọc)", "Chỉ các bản lounge/instrumental nhẹ; loại tango/latin hoặc nhịp nổi bật."),
}


def main() -> None:
    source = load_workbook(SOURCE, read_only=True, data_only=True)
    audit_rows = json.loads(AUDIT.read_text(encoding="utf-8")) if AUDIT.exists() else []
    audio_presence = {}
    for row in audit_rows:
        audio_presence[(row["sheet"], row["no"])] = "; ".join(row.get("existing", []))

    wb = load_workbook(SOURCE)
    if "Review_3_Channels" in wb.sheetnames:
        del wb["Review_3_Channels"]
    ws = wb.create_sheet("Review_3_Channels")
    headers = ["Pool", "Thư mục thư viện", "STT", "Tên bài", "Nghệ sĩ", "Fit Tier gốc", "Priority gốc",
               "Best Use gốc", "Kênh đề xuất", "Nhận định sơ bộ", "File audio cục bộ", "Nghe kiểm tra",
               "Tempo / nhịp", "Lời / instrumental", "Quyết định", "Ghi chú nghe", "Link tra cứu"]
    ws.append(headers)
    for sheet in source.worksheets:
        if sheet.title == "README" or sheet.title not in POOLS:
            continue
        folder, channels, rationale = POOLS[sheet.title]
        rows = list(sheet.values)
        for values in rows[4:]:
            if not values or not isinstance(values[0], (int, float)) or not values[1]:
                continue
            number, title, artist = int(values[0]), values[1], values[2]
            ws.append([sheet.title, folder, number, title, artist, values[3], values[4], values[5], channels,
                       rationale, audio_presence.get((sheet.title, number), "Chưa đối chiếu"), "Chưa nghe",
                       "Chưa xác định", "Chưa xác định", "Chờ nghe", "", values[12] if len(values) > 12 else ""])

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    widths = [20, 18, 8, 42, 25, 18, 14, 42, 38, 62, 70, 16, 18, 22, 16, 42, 56]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="254E70")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (8, 9, 10, 11, 16, 17))
        row[11].fill = PatternFill("solid", fgColor="FFF2CC")
        row[14].fill = PatternFill("solid", fgColor="FCE4D6")
    table = Table(displayName="MusicReview", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                          showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
    ws.add_table(table)

    guide = wb.create_sheet("Huong_dan")
    guide.append(["MỤC TIÊU", "Rà soát 600 placement nhạc cho HEARTHLUNE, SYLVANREST và ASTRALOOM."])
    guide.append(["PHẠM VI", "Danh sách lấy từ YouTube_Library_600_Tracks.xlsx; một bài có thể xuất hiện ở nhiều pool."])
    guide.append(["GIỚI HẠN", "Fit Tier/Best Use là metadata gốc. Tempo, mức năng lượng, lời hát và chất lượng loop chưa được nghe xác nhận."])
    guide.append(["CÁCH DÙNG", "Lọc theo Kênh đề xuất; nghe file cục bộ hoặc tra YouTube Audio Library; điền Tempo/nhịp, Lời/instrumental và Quyết định."])
    guide.append(["QUYẾT ĐỊNH", "Chỉ giữ khi không lời, âm lượng/nhịp nhẹ và hợp cảnh. Đánh dấu Loại nếu beat mạnh, vocal, tango/latin nổi bật hoặc không hợp định vị."])
    guide.append(["KHÔNG TỰ XÓA", "Workbook chỉ đề xuất. Không có bản nhạc nào bị xóa khỏi ổ đĩa trong quá trình tạo bảng."])
    guide.append(["KÊNH 3", "Ưu tiên ambient/drone/white-noise; nhạc có tiết tấu rõ chỉ giữ nếu nghe thực tế xác nhận rất nhẹ, nếu không thì loại."])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 115
    for row in guide.iter_rows():
        row[0].font = Font(bold=True, color="FFFFFF")
        row[0].fill = PatternFill("solid", fgColor="254E70")
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(DEST)
    print(f"Created {DEST.name} with {ws.max_row - 1} placements")


if __name__ == "__main__":
    main()
